"""Files in, located chunks out.

Two rules this module exists to enforce:

1. A chunk's source location is attached *inside* the parse loop, from the thing
   that actually produced the text. It is never reconstructed afterwards by
   searching for a string, because that guess is what makes citations untrustworthy.

2. chunk_id hashes the document's content hash plus the canonical location, and
   NEVER the text. OCR output varies between Tesseract versions and settings, so
   hashing text would make IDs unstable on exactly the scanned documents where
   stability matters most.
"""
from __future__ import annotations

import hashlib
from pathlib import Path

import pymupdf
import pytesseract
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from PIL import Image

from docint.config import (
    MIN_TEXT_LAYER_CHARS,
    RASTER_PAGE_COVERAGE,
    OCR_RENDER_DPI,
    access_tag_for,
)
from docint.models import Chunk, Document, FileType, OcrTelemetry, SourceLocation


class UnsupportedFileType(Exception):
    """Raised for a file no parser claims - surfaced as an ingest-stage failure."""


def content_hash(path: Path) -> str:
    h = hashlib.sha256()
    with open(path, "rb") as fh:
        for block in iter(lambda: fh.read(1 << 20), b""):
            h.update(block)
    return h.hexdigest()


def make_chunk_id(doc_content_hash: str, location: SourceLocation) -> str:
    return hashlib.sha256(
        f"{doc_content_hash}|{location.canonical()}".encode()
    ).hexdigest()[:16]


# --------------------------------------------------------------------------- #
# OCR
# --------------------------------------------------------------------------- #

def ocr_page(page: pymupdf.Page) -> tuple[str, float]:
    """Rasterise and OCR one page. Returns (text, mean word confidence 0-100).

    Text is rebuilt line by line from Tesseract's block/paragraph/line indices
    rather than space-joining every word, because collapsing a form to one long
    line measurably degrades downstream field extraction.
    """
    pix = page.get_pixmap(dpi=OCR_RENDER_DPI)
    img = Image.frombytes("RGB", (pix.width, pix.height), pix.samples)
    data = pytesseract.image_to_data(img, output_type=pytesseract.Output.DICT)

    lines: dict[tuple[int, int, int], list[str]] = {}
    confidences: list[float] = []
    for i, word in enumerate(data["text"]):
        if not word.strip():
            continue
        conf = float(data["conf"][i])
        if conf < 0:                       # Tesseract's "no confidence" sentinel
            continue
        key = (data["block_num"][i], data["par_num"][i], data["line_num"][i])
        lines.setdefault(key, []).append(word)
        confidences.append(conf)

    text = "\n".join(" ".join(words) for _, words in sorted(lines.items()))
    mean_conf = sum(confidences) / len(confidences) if confidences else 0.0
    return text, mean_conf


# --------------------------------------------------------------------------- #
# PDF
# --------------------------------------------------------------------------- #

def raster_coverage(page) -> float:
    """Fraction of the page covered by its largest raster image.

    The largest rather than the sum, because overlapping images would otherwise
    total more than the page. This is how a scan is recognised even when it carries
    an inherited text layer - see RASTER_PAGE_COVERAGE in config.
    """
    images = page.get_images(full=True)
    if not images:
        return 0.0
    area = page.rect.width * page.rect.height
    if area <= 0:
        return 0.0
    largest = 0.0
    for image in images:
        try:
            bbox = page.get_image_bbox(image)
        except (ValueError, RuntimeError):     # malformed or unplaceable image
            continue
        largest = max(largest, abs(bbox.width * bbox.height))
    return largest / area


def parse_pdf(path: Path, doc_hash: str, access_tag: str) -> tuple[list[Chunk], FileType, float | None, int]:
    """One chunk per page. Each page independently takes the text-layer or OCR path.

    Real-world PDFs are frequently mixed - a born-digital cover page in front of
    scanned attachments - so the decision is per page, not per file. The document
    is labelled `pdf_scanned` if ANY page needed OCR, since that is what determines
    whether its content carries recognition risk.
    """
    pdf = pymupdf.open(path)
    chunks: list[Chunk] = []
    ocr_confidences: list[float] = []
    used_ocr = False

    try:
        for page_index, page in enumerate(pdf):
            page_number = page_index + 1
            text = page.get_text().strip()
            confidence: float | None = None

            # Two ways a page earns OCR: it has almost no text, or it is a raster
            # scan that merely came with text attached. In the second case the
            # inherited layer is DISCARDED rather than merged - a value this system
            # asserts should be one it measured the reading of, and text of unknown
            # provenance carries no confidence to attach to a citation.
            if len(text) < MIN_TEXT_LAYER_CHARS or raster_coverage(page) >= RASTER_PAGE_COVERAGE:
                text, confidence = ocr_page(page)
                used_ocr = True
                ocr_confidences.append(confidence)

            if not text.strip():
                continue

            location = SourceLocation(kind="pdf_page", page=page_number)
            chunks.append(
                Chunk(
                    chunk_id=make_chunk_id(doc_hash, location),
                    document_id=f"doc_{doc_hash[:12]}",
                    document_type="unknown",        # stamped after classification
                    file_type="pdf_scanned" if used_ocr else "pdf_digital",
                    filename=path.name,
                    access_tag=access_tag,
                    text=text,
                    source_location=location,
                    ocr_confidence=confidence,
                )
            )
        page_count = pdf.page_count
    finally:
        pdf.close()

    file_type: FileType = "pdf_scanned" if used_ocr else "pdf_digital"
    mean_conf = sum(ocr_confidences) / len(ocr_confidences) if ocr_confidences else None
    return chunks, file_type, mean_conf, page_count


# --------------------------------------------------------------------------- #
# XLSX
# --------------------------------------------------------------------------- #

def _row_blocks(ws) -> list[tuple[int, int]]:
    """Maximal runs of consecutive non-empty rows. Blank rows are the separator."""
    populated = [
        r for r in range(1, ws.max_row + 1)
        if any(ws.cell(row=r, column=c).value is not None for c in range(1, ws.max_column + 1))
    ]
    blocks: list[tuple[int, int]] = []
    for row in populated:
        if blocks and row == blocks[-1][1] + 1:
            blocks[-1] = (blocks[-1][0], row)
        else:
            blocks.append((row, row))
    return blocks


def parse_xlsx(path: Path, doc_hash: str, access_tag: str) -> tuple[list[Chunk], FileType, None, int]:
    """One chunk per contiguous row block, located by sheet and cell range.

    A block is the natural unit here: a labelled terms section and a line-item
    table are separate things a citation should be able to point at individually.
    """
    wb = load_workbook(path, data_only=True)
    chunks: list[Chunk] = []

    for ws in wb.worksheets:
        for start_row, end_row in _row_blocks(ws):
            cols = [
                c for c in range(1, ws.max_column + 1)
                if any(ws.cell(row=r, column=c).value is not None
                       for r in range(start_row, end_row + 1))
            ]
            if not cols:
                continue
            first_col, last_col = min(cols), max(cols)

            lines = []
            for r in range(start_row, end_row + 1):
                values = [
                    str(ws.cell(row=r, column=c).value)
                    for c in range(first_col, last_col + 1)
                    if ws.cell(row=r, column=c).value is not None
                ]
                if values:
                    lines.append(" | ".join(values))
            text = "\n".join(lines)
            if not text.strip():
                continue

            location = SourceLocation(
                kind="xlsx_range",
                sheet=ws.title,
                cell_range=f"{get_column_letter(first_col)}{start_row}:"
                           f"{get_column_letter(last_col)}{end_row}",
            )
            chunks.append(
                Chunk(
                    chunk_id=make_chunk_id(doc_hash, location),
                    document_id=f"doc_{doc_hash[:12]}",
                    document_type="unknown",
                    file_type="xlsx",
                    filename=path.name,
                    access_tag=access_tag,
                    text=text,
                    source_location=location,
                    ocr_confidence=None,
                )
            )
    sheet_count = len(wb.worksheets)
    wb.close()
    return chunks, "xlsx", None, sheet_count


# --------------------------------------------------------------------------- #

def parse(path: Path) -> tuple[Document, list[Chunk]]:
    """Parse one file into a Document and its located Chunks."""
    path = Path(path)
    doc_hash = content_hash(path)
    access_tag = access_tag_for(path)
    suffix = path.suffix.lower()

    if suffix == ".pdf":
        chunks, file_type, mean_conf, unit_count = parse_pdf(path, doc_hash, access_tag)
    elif suffix in (".xlsx", ".xlsm"):
        chunks, file_type, mean_conf, unit_count = parse_xlsx(path, doc_hash, access_tag)
    else:
        raise UnsupportedFileType(f"no parser for {path.suffix!r} ({path.name})")

    document = Document(
        document_id=f"doc_{doc_hash[:12]}",
        filename=path.name,
        file_type=file_type,
        content_hash=doc_hash,
        access_tag=access_tag,
        page_count=unit_count,
        ocr=OcrTelemetry(tesseract_confidence=mean_conf),
    )
    return document, chunks


def finalize_chunks(chunks: list[Chunk], document: Document) -> list[Chunk]:
    """Stamp post-classification document facts onto each chunk.

    Chunks are built during parsing, before the document type is known. This
    cannot change chunk_id, because identity hashes content and location only.
    """
    return [c.model_copy(update={"document_type": document.document_type}) for c in chunks]
