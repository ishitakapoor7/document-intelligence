"""Author the three demo documents deterministically, plus the L1/L2/L3 scan ladder.

Everything here is synthetic. The numbers are chosen so the cross-document question
has a hand-checkable answer:

    invoice : 80 units @ $156.00 = $12,480.00
    vendor  : contracted rate    = $150.00/unit   -> 80 x $6.00 = $480.00 overcharge
    PO      : committed          = $12,000.00     -> $12,480 - $12,000 = $480.00 over

Both overages are exactly $480.00.

Run:  python scripts/build_corpus.py
"""
from __future__ import annotations

import io
import random
import shutil
from pathlib import Path

import pymupdf
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from PIL import Image, ImageFilter
from reportlab.lib.pagesizes import LETTER
from reportlab.lib.units import inch
from reportlab.pdfgen import canvas

ROOT = Path(__file__).resolve().parent.parent
CORPUS = ROOT / "corpus"
SCANS = ROOT / "eval" / "scans"
BUILD = ROOT / "runs" / "_build"

SEED = 20260918  # fixed so degradation is reproducible

# --------------------------------------------------------------------------- #
# 1. Digital invoice  (real text layer -> parsed without OCR)
# --------------------------------------------------------------------------- #

def build_invoice(path: Path) -> None:
    c = canvas.Canvas(str(path), pagesize=LETTER)
    w, h = LETTER
    y = h - 1.0 * inch

    c.setFont("Helvetica-Bold", 20)
    c.drawString(1 * inch, y, "INVOICE")
    c.setFont("Helvetica", 10)
    c.drawRightString(w - 1 * inch, y, "Acme Industrial Supply Co.")
    c.drawRightString(w - 1 * inch, y - 14, "4417 Foundry Road, Akron, OH 44311")
    c.drawRightString(w - 1 * inch, y - 28, "Vendor ID: V-1042")

    y -= 70
    c.setFont("Helvetica-Bold", 11)
    c.drawString(1 * inch, y, "Bill To:")
    c.setFont("Helvetica", 10)
    c.drawString(1 * inch, y - 15, "Northwind Manufacturing LLC")
    c.drawString(1 * inch, y - 29, "2100 Lakeshore Drive, Milwaukee, WI 53202")

    rows = [
        ("Invoice Number", "INV-2026-0117"),
        ("Invoice Date", "2026-03-14"),
        ("PO Reference", "PO-2026-0043"),
        ("Payment Terms", "Net 30"),
        ("Currency", "USD"),
    ]
    yy = y
    for label, value in rows:
        c.setFont("Helvetica-Bold", 10)
        c.drawRightString(w - 2.4 * inch, yy, f"{label}:")
        c.setFont("Helvetica", 10)
        c.drawRightString(w - 1 * inch, yy, value)
        yy -= 15

    y -= 80
    c.setFont("Helvetica-Bold", 10)
    for x, t in ((1.0, "Description"), (4.6, "Qty"), (5.5, "Unit Price"), (6.9, "Amount")):
        c.drawString(x * inch, y, t)
    c.line(1 * inch, y - 5, w - 1 * inch, y - 5)

    y -= 22
    c.setFont("Helvetica", 10)
    c.drawString(1.0 * inch, y, "Grade-A Hydraulic Coupling, 1/2 in. NPT")
    c.drawString(4.6 * inch, y, "80")
    c.drawRightString(6.6 * inch, y, "$156.00")
    c.drawRightString(w - 1 * inch, y, "$12,480.00")

    y -= 40
    c.line(4.6 * inch, y + 12, w - 1 * inch, y + 12)
    for label, value, bold in (
        ("Subtotal", "$12,480.00", False),
        ("Tax (0%)", "$0.00", False),
        ("Total Due", "$12,480.00", True),
    ):
        c.setFont("Helvetica-Bold" if bold else "Helvetica", 11 if bold else 10)
        c.drawRightString(6.6 * inch, y, f"{label}:")
        c.drawRightString(w - 1 * inch, y, value)
        y -= 17

    c.setFont("Helvetica-Oblique", 8)
    c.drawString(1 * inch, 0.8 * inch,
                 "Remit within 30 days. Reference the invoice number on all payments.")
    c.showPage()
    c.save()


# --------------------------------------------------------------------------- #
# 2. Purchase order  (rendered, then rasterised -> NO text layer -> OCR path)
# --------------------------------------------------------------------------- #

def build_po_source(path: Path) -> None:
    c = canvas.Canvas(str(path), pagesize=LETTER)
    w, h = LETTER
    y = h - 1.0 * inch

    c.setFont("Helvetica-Bold", 20)
    c.drawString(1 * inch, y, "PURCHASE ORDER")
    c.setFont("Helvetica", 10)
    c.drawRightString(w - 1 * inch, y, "Northwind Manufacturing LLC")
    c.drawRightString(w - 1 * inch, y - 14, "2100 Lakeshore Drive, Milwaukee, WI 53202")

    y -= 64
    rows = [
        ("PO Number", "PO-2026-0043"),
        ("Issue Date", "2026-02-28"),
        ("Buyer Entity", "Northwind Manufacturing LLC"),
        ("Supplier", "Acme Industrial Supply Co."),
        ("Currency", "USD"),
    ]
    for label, value in rows:
        c.setFont("Helvetica-Bold", 11)
        c.drawString(1 * inch, y, f"{label}:")
        c.setFont("Helvetica", 11)
        c.drawString(2.9 * inch, y, value)
        y -= 19

    y -= 24
    c.setFont("Helvetica-Bold", 11)
    for x, t in ((1.0, "Line"), (1.7, "Description"), (5.3, "Qty"), (6.2, "Not-To-Exceed")):
        c.drawString(x * inch, y, t)
    c.line(1 * inch, y - 6, w - 1 * inch, y - 6)

    y -= 26
    c.setFont("Helvetica", 11)
    c.drawString(1.0 * inch, y, "1")
    c.drawString(1.7 * inch, y, "Grade-A Hydraulic Coupling, 1/2 in. NPT")
    c.drawString(5.3 * inch, y, "80")
    c.drawRightString(w - 1 * inch, y, "$12,000.00")

    y -= 46
    c.setFont("Helvetica-Bold", 13)
    c.drawRightString(6.4 * inch, y, "Total Committed Amount:")
    c.drawRightString(w - 1 * inch, y, "$12,000.00")

    y -= 46
    c.setFont("Helvetica", 9)
    c.drawString(1 * inch, y, "Authorised by: D. Okafor, Procurement Manager")
    c.drawString(1 * inch, y - 13,
                 "Invoices exceeding the committed amount require a written change order.")
    c.showPage()
    c.save()


# Degradation ladder. L1 is the readable copy that also ships in corpus/.
#
# MEASURED THROUGH THE REAL PIPELINE (degrade -> image-only PDF -> parse() -> OCR),
# not through a probe. This matters: an earlier tuning pass rasterised the *source*
# PDF at the target DPI and OCR'd that directly, reporting 76.9 for L2. The pipeline
# instead re-renders the already-degraded image-only PDF back up to 300 DPI, and that
# upscaling adds interpolation damage the probe never saw - the true figure was 41.3.
# Always measure the path the system actually takes.
#
#   L1  ocr 95.1  - all five gold fields readable
#   L2  ocr 62.4  - all five still readable, but below the field-flag threshold,
#                   so every extracted value is surfaced as LOW-CONFIDENCE
#   L3  ocr 43.4  - vendor_name and po_number lost; below MIN_OCR_CONFIDENCE,
#                   so extraction is skipped and the document goes to human review
#
# Thresholds follow from these observations: MIN_OCR_CONFIDENCE=55 sits between L2
# and L3; LOW_CONFIDENCE_FIELD=75 sits between L1 and L2.
LEVELS = {
    "L1_clean":    dict(dpi=300, blur=0.0, noise=0, jpeg=95, rotate=0.0),
    "L2_medium":   dict(dpi=200, blur=1.0, noise=4, jpeg=35, rotate=0.6),
    "L3_degraded": dict(dpi=150, blur=1.2, noise=5, jpeg=30, rotate=0.7),
}


def rasterise(src_pdf: Path, dpi: int) -> Image.Image:
    doc = pymupdf.open(src_pdf)
    pix = doc[0].get_pixmap(dpi=dpi)
    img = Image.frombytes("RGB", (pix.width, pix.height), pix.samples)
    doc.close()
    return img


def degrade(img: Image.Image, *, blur: float, noise: int, jpeg: int,
            rotate: float, seed: int) -> Image.Image:
    rng = random.Random(seed)
    if rotate:
        img = img.rotate(rotate, resample=Image.BICUBIC, fillcolor=(255, 255, 255))
    if blur:
        img = img.filter(ImageFilter.GaussianBlur(blur))
    if noise:
        px = img.load()
        for _ in range(int(img.width * img.height * noise / 1000)):
            x, y = rng.randrange(img.width), rng.randrange(img.height)
            v = rng.randint(0, 255)
            px[x, y] = (v, v, v)
    buf = io.BytesIO()
    img.convert("L").save(buf, format="JPEG", quality=jpeg)
    buf.seek(0)
    return Image.open(buf).convert("RGB")


def image_to_pdf(img: Image.Image, out: Path) -> None:
    """Wrap a bitmap as a single-page PDF with NO text layer -> forces the OCR path.

    Embedded as grayscale JPEG: a lossless PNG embed made L1 a 25 MB file, which is
    not something to put in a repo a reviewer clones.
    """
    buf = io.BytesIO()
    img.convert("L").save(buf, format="JPEG", quality=90, optimize=True)
    doc = pymupdf.open()
    page = doc.new_page(width=612, height=792)          # US Letter, points
    page.insert_image(pymupdf.Rect(0, 0, 612, 792), stream=buf.getvalue())
    doc.save(out)
    doc.close()


# --------------------------------------------------------------------------- #
# 3. Vendor record workbook
# --------------------------------------------------------------------------- #

def build_vendor_records(path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Vendors"
    bold = Font(bold=True)

    ws["A1"] = "VENDOR MASTER RECORD"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = "Northwind Manufacturing LLC - Procurement"

    terms = [
        ("Vendor ID", "V-1042"),
        ("Vendor Name", "Acme Industrial Supply Co."),
        ("Contracted Unit Price", 150.00),
        ("Contract Effective Date", "2026-01-01"),
        ("Payment Terms", "Net 30"),
        ("Contract Status", "Active"),
    ]
    for i, (label, value) in enumerate(terms, start=4):
        ws.cell(row=i, column=1, value=label).font = bold
        cell = ws.cell(row=i, column=2, value=value)
        if isinstance(value, float):
            cell.number_format = '"$"#,##0.00'

    ws["A12"] = "CONTRACTED LINE ITEMS"
    ws["A12"].font = bold
    headers = ["SKU", "Description", "UoM", "Contracted Unit Price", "Min Qty"]
    for j, head in enumerate(headers, start=1):
        c = ws.cell(row=13, column=j, value=head)
        c.font = bold
        c.alignment = Alignment(horizontal="center")

    items = [
        ("AC-4410", "Grade-A Hydraulic Coupling, 1/2 in. NPT", "EA", 150.00, 25),
        ("AC-4412", "Grade-A Hydraulic Coupling, 3/4 in. NPT", "EA", 168.50, 25),
        ("AC-7701", "PTFE Thread Sealant Tape, 12mm x 20m",     "RL",   4.75, 100),
        ("AC-9120", "Stainless Compression Fitting, 1/2 in.",   "EA",  38.20, 50),
    ]
    for i, row in enumerate(items, start=14):
        for j, value in enumerate(row, start=1):
            c = ws.cell(row=i, column=j, value=value)
            if j == 4:
                c.number_format = '"$"#,##0.00'

    for col, width in zip("ABCDE", (26, 42, 8, 22, 10)):
        ws.column_dimensions[col].width = width
    wb.save(path)


# --------------------------------------------------------------------------- #

def main() -> None:
    for d in (CORPUS / "procurement", CORPUS / "general", SCANS, BUILD):
        d.mkdir(parents=True, exist_ok=True)

    build_invoice(CORPUS / "procurement" / "invoice_acme_001.pdf")
    print(f"  wrote {CORPUS / 'procurement' / 'invoice_acme_001.pdf'}")

    src = BUILD / "po_source.pdf"
    build_po_source(src)
    for name, cfg in LEVELS.items():
        img = degrade(rasterise(src, cfg["dpi"]), blur=cfg["blur"], noise=cfg["noise"],
                      jpeg=cfg["jpeg"], rotate=cfg["rotate"], seed=SEED)
        out = SCANS / f"po_acme_001_{name}.pdf"
        image_to_pdf(img, out)
        print(f"  wrote {out}  ({cfg['dpi']} dpi)")

    # the demo copy IS the clean rung of the ladder - same bytes, no third render
    shutil.copyfile(SCANS / "po_acme_001_L1_clean.pdf", CORPUS / "procurement" / "po_acme_001.pdf")
    print(f"  wrote {CORPUS / 'procurement' / 'po_acme_001.pdf'}  (copy of L1_clean)")

    build_vendor_records(CORPUS / "general" / "vendor_records.xlsx")
    print(f"  wrote {CORPUS / 'general' / 'vendor_records.xlsx'}")


if __name__ == "__main__":
    main()
